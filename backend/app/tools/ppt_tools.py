"""可编辑的中国联通风格 PowerPoint 生成工具。"""
from __future__ import annotations

import uuid
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from langchain.tools import tool
from langgraph.config import get_config
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app.paths import WORKSPACE_DATA, ROOT as BACKEND_ROOT

RED = RGBColor(0xE6, 0x00, 0x27)
BLACK = RGBColor(0x16, 0x16, 0x16)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"
SLIDE_W = 13.333333
SLIDE_H = 7.5
LOGO = BACKEND_ROOT / "skills" / "china-unicom-ppt" / "assets" / "logo.png"


def _text(slide, text: str, x: float, y: float, w: float, h: float, *,
          size: int, bold: bool = False, color=BLACK, align=PP_ALIGN.LEFT,
          valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.08
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _rect(slide, x: float, y: float, w: float, h: float, color=RED):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _base_slide(prs: Presentation, title: str, page: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    _text(slide, title, .52, .24, 10.6, .62, size=23, bold=True)
    slide.shapes.add_picture(str(LOGO), Inches(11.15), Inches(.22), width=Inches(1.65))
    _rect(slide, 0, 1.0, SLIDE_W, .035)
    _text(slide, f"{page}", 12.45, 6.94, .45, .22, size=9, color=GRAY,
          align=PP_ALIGN.RIGHT)
    return slide


def _build_presentation(title: str, subtitle: str, points: list[str], output: Path) -> None:
    if not LOGO.is_file():
        raise RuntimeError("中国联通 Logo 资源缺失，无法生成品牌演示文稿")
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    # 封面：保留参考模板的白底、左侧红色竖线与右上角品牌标识。
    cover = prs.slides.add_slide(prs.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = WHITE
    cover.shapes.add_picture(str(LOGO), Inches(8.85), Inches(.48), width=Inches(3.2))
    _rect(cover, 1.62, 2.12, .065, 3.18)
    _text(cover, title, 2.15, 2.28, 9.7, 1.28, size=34, bold=True)
    _text(cover, subtitle or "中国联通风格演示稿", 2.18, 3.92, 9.6, .66,
          size=17, color=GRAY)
    today = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y年%m月%d日")
    _text(cover, today, 2.18, 4.78, 8.2, .5, size=14, color=GRAY)
    _text(cover, "中国联通 · HRBP", 2.18, 6.82, 4.0, .24, size=9, color=GRAY)

    # 每页最多三个观点，避免长文本挤在一页。
    batches = [points[i:i + 3] for i in range(0, len(points), 3)]
    for page_index, batch in enumerate(batches, start=2):
        slide = _base_slide(prs, "关键要点" if page_index == 2 else "关键要点（续）", page_index)
        row_top = 1.40
        row_height = 1.70
        for item_index, point in enumerate(batch):
            y = row_top + item_index * row_height
            number = (page_index - 2) * 3 + item_index + 1
            _text(slide, f"{number:02d}", .62, y + .08, .64, .60, size=19,
                  bold=True, color=RED, align=PP_ALIGN.CENTER)
            _rect(slide, 1.43, y + .10, .035, 1.18)
            _text(slide, point, 1.73, y, 10.65, 1.42, size=18, color=BLACK)

    closing = prs.slides.add_slide(prs.slide_layouts[6])
    closing.background.fill.solid()
    closing.background.fill.fore_color.rgb = WHITE
    closing.shapes.add_picture(str(LOGO), Inches(5.02), Inches(1.28), width=Inches(3.28))
    _text(closing, "感谢聆听", 2.25, 3.05, 8.85, .9, size=34, bold=True,
          align=PP_ALIGN.CENTER)
    _text(closing, "守正创新  行稳致远", 3.0, 4.15, 7.35, .5, size=17,
          color=GRAY, align=PP_ALIGN.CENTER)
    prs.save(output)


@tool
def generate_china_unicom_ppt(title: str, key_points: str,
                              subtitle: str = "") -> str:
    """【生成中国联通 PPT】生成可编辑的 16:9 中国联通品牌 PowerPoint 文件。

    适用于中国联通公司介绍、项目方案、工作汇报或技术分享。按用户给出的
    主题与要点生成封面、要点页和结束页；每行一个要点，最多 8 行。请勿补造
    未提供的经营数字、政策事实、案例或来源。文件会出现在会话产物卡片，
    并可在线预览。

    参数:
        title: 演示文稿标题。
        key_points: 关键内容，每行一个观点，建议 3 至 6 行。
        subtitle: 副标题或汇报对象（可选）。
    """
    title = (title or "").strip()
    subtitle = (subtitle or "").strip()
    points = [line.strip().lstrip("•*- ").strip()
              for line in (key_points or "").splitlines() if line.strip()]
    if not title or len(title) > 64:
        raise ValueError("PPT 标题不能为空且最多 64 个字符")
    if not 1 <= len(points) <= 8:
        raise ValueError("请提供 1 至 8 条要点，每行一条")
    if any(len(point) > 110 for point in points):
        raise ValueError("每条要点请压缩到 110 个字符以内，避免版面溢出")
    if len(subtitle) > 100:
        raise ValueError("副标题最多 100 个字符")

    try:
        cfg = get_config() or {}
        user_id = (cfg.get("configurable") or {}).get("user_id")
    except Exception as exc:
        raise RuntimeError(f"无法获取运行时用户身份，拒绝生成文件：{exc}") from exc
    if not isinstance(user_id, str) or not user_id.strip():
        raise RuntimeError("运行时缺少 user_id，拒绝将文件写入共享目录")
    user_key = user_id.strip()
    if (not user_key or user_key in {".", ".."} or user_key.startswith(".")
            or "/" in user_key or "\\" in user_key or "\x00" in user_key):
        raise RuntimeError("运行时 user_id 无效，无法创建用户专属目录")

    user_dir = (WORKSPACE_DATA / user_key).resolve()
    root = WORKSPACE_DATA.resolve()
    if root not in user_dir.parents:
        raise RuntimeError("用户工作区路径越界，已拒绝生成")
    user_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y%m%d_%H%M%S")
    filename = f"China_Unicom_{stamp}_{uuid.uuid4().hex[:8]}.pptx"
    output = user_dir / filename
    _build_presentation(title, subtitle, points, output)
    size = output.stat().st_size
    relative = output.relative_to(root).as_posix()
    slide_count = (len(points) + 2) // 3 + 2
    return (f"已生成中国联通风格 PowerPoint（{slide_count} 页，{size:,} 字节），"
            f"支持在线预览：/workspace/data/{relative}")
