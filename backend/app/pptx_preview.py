"""把 PowerPoint 文字与图片渲染为静态 HTML，供受控 iframe 在线预览。"""
from __future__ import annotations

import base64
import html
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

MAX_PPTX_SIZE = 40 * 1024 * 1024
MAX_UNPACKED_SIZE = 160 * 1024 * 1024
MAX_INLINE_IMAGES_SIZE = 20 * 1024 * 1024
MAX_SLIDES = 100
MAX_SHAPES = 2500
ALLOWED_IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/bmp", "image/webp"}


def _check_package(path: Path) -> None:
    if path.stat().st_size > MAX_PPTX_SIZE:
        raise ValueError("PPTX 超过 40 MB，暂不支持在线预览，请下载后查看")
    try:
        with zipfile.ZipFile(path) as package:
            if sum(item.file_size for item in package.infolist()) > MAX_UNPACKED_SIZE:
                raise ValueError("PPTX 解压内容过大，暂不支持在线预览")
            if "ppt/presentation.xml" not in package.namelist():
                raise ValueError("文件不是有效的 PPTX 文档")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("文件不是有效的 PPTX 文档") from exc


def _rgb(value, default: str) -> str:
    try:
        raw = str(value)
        if re.fullmatch(r"[0-9A-Fa-f]{6}", raw):
            return f"#{raw}"
    except Exception:
        pass
    return default


def _shape_position(shape, slide_width: int, slide_height: int) -> str:
    left = max(0, min(100, shape.left / slide_width * 100))
    top = max(0, min(100, shape.top / slide_height * 100))
    width = max(0, min(100, shape.width / slide_width * 100))
    height = max(0, min(100, shape.height / slide_height * 100))
    return (f"left:{left:.5f}%;top:{top:.5f}%;width:{width:.5f}%;"
            f"height:{height:.5f}%;")


def _shape_fill(shape) -> str | None:
    try:
        if shape.fill.type is None:
            return None
        if shape.fill.type == MSO_FILL.SOLID:
            return _rgb(shape.fill.fore_color.rgb, "#ffffff")
    except Exception:
        return None
    return None


def _text_shape_html(shape, slide_width: int, slide_height: int) -> str:
    frame = shape.text_frame
    margin_left = max(0, frame.margin_left / shape.width * 100) if shape.width else 0
    margin_right = max(0, frame.margin_right / shape.width * 100) if shape.width else 0
    margin_top = max(0, frame.margin_top / shape.height * 100) if shape.height else 0
    margin_bottom = max(0, frame.margin_bottom / shape.height * 100) if shape.height else 0
    paragraphs = []
    for paragraph in frame.paragraphs:
        parts = []
        for run in paragraph.runs:
            content = html.escape(run.text or "").replace("\n", "<br>").replace("\t", "&emsp;")
            if not content:
                continue
            font = run.font
            styles = []
            if font.size and 4 <= font.size.pt <= 96:
                cqw = font.size.pt * 1.3333 / 12.8
                styles.append(f"font-size:clamp(6px,{cqw:.4f}cqw,96px)")
            if font.color and font.color.type is not None:
                try:
                    styles.append(f"color:{_rgb(font.color.rgb, '#161616')}")
                except Exception:
                    pass
            if font.bold:
                styles.append("font-weight:700")
            if font.italic:
                styles.append("font-style:italic")
            if styles:
                content = f'<span style="{html.escape(";".join(styles), quote=True)}">{content}</span>'
            parts.append(content)
        if not parts and paragraph.text:
            parts.append(html.escape(paragraph.text).replace("\n", "<br>"))
        align = "left"
        if paragraph.alignment == 2:
            align = "center"
        elif paragraph.alignment == 3:
            align = "right"
        elif paragraph.alignment == 4:
            align = "justify"
        paragraphs.append(f'<div style="text-align:{align}">{"".join(parts) or "&nbsp;"}</div>')
    content = "".join(paragraphs)
    return (f'<div class="item text" style="{_shape_position(shape, slide_width, slide_height)}'
            f'padding:{margin_top:.3f}% {margin_right:.3f}% {margin_bottom:.3f}% {margin_left:.3f}%">'
            f'{content}</div>')


def _render_shape(shape, slide_width: int, slide_height: int,
                  image_budget: list[int]) -> str:
    pos = _shape_position(shape, slide_width, slide_height)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        image = shape.image
        content_type = image.content_type
        blob = image.blob
        if (content_type not in ALLOWED_IMAGE_TYPES or
                image_budget[0] + len(blob) > MAX_INLINE_IMAGES_SIZE):
            return ""
        image_budget[0] += len(blob)
        encoded = base64.b64encode(blob).decode("ascii")
        return (f'<img class="item picture" style="{pos}" alt="" '
                f'src="data:{content_type};base64,{encoded}">')
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return _text_shape_html(shape, slide_width, slide_height)
    fill = _shape_fill(shape)
    if fill:
        return f'<div class="item shape" style="{pos}background:{fill}"></div>'
    return ""


def render_pptx_html(path: Path) -> str:
    """Render a bounded, script-free slide overview from an authorized PPTX."""
    _check_package(path)
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ValueError("PPTX 无法解析，请检查文件是否损坏") from exc
    if not 1 <= len(prs.slides) <= MAX_SLIDES:
        raise ValueError("PPTX 页数需在 1 至 100 页之间才能在线预览")
    shape_count = sum(len(slide.shapes) for slide in prs.slides)
    if shape_count > MAX_SHAPES:
        raise ValueError("PPTX 页面元素过多，暂不支持在线预览")

    width, height = prs.slide_width, prs.slide_height
    aspect_ratio = width / height
    budget = [0]
    rendered_slides = []
    for index, slide in enumerate(prs.slides, start=1):
        elements = "".join(_render_shape(shape, width, height, budget)
                            for shape in slide.shapes)
        rendered_slides.append(
            f'<section class="slide"><div class="slide-stage" style="aspect-ratio:{aspect_ratio:.6f}">{elements}</div>'
            f'<div class="slide-number">{index} / {len(prs.slides)}</div></section>')
    return """<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>PowerPoint 在线预览</title><style>
*{box-sizing:border-box}body{margin:0;padding:18px;background:#f1f3f5;color:#161616;font-family:Arial,"Microsoft YaHei",sans-serif}
.slide{container-type:inline-size;max-width:1280px;margin:0 auto 22px;background:#fff;box-shadow:0 2px 8px #0002}
.slide-stage{position:relative;width:100%;overflow:hidden;background:#fff}
.item{position:absolute}.text{font-size:clamp(6px,2cqw,32px);line-height:1.16;white-space:pre-wrap;overflow:hidden;overflow-wrap:anywhere}
.picture{object-fit:contain}.shape{border:0}.slide-number{padding:5px 12px;text-align:right;color:#707070;font-size:12px;background:#fff}
@media(max-width:600px){body{padding:8px}.slide{margin-bottom:12px}}
</style></head><body>""" + "".join(rendered_slides) + "</body></html>"
