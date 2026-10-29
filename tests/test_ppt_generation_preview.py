from pathlib import Path

from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from app import auth, conversations
from app.catalog import db as catalog_db
from app.catalog import seeds
from app.skill_runtime import skill_snapshot
import app.paths as app_paths
import app.routes.workspace as workspace
from app.tools import ppt_tools


def test_hrbp_backfill_binds_default_skill_and_ppt_tool(monkeypatch, tmp_path):
    monkeypatch.setattr(catalog_db, "DB", tmp_path / "catalog.db")
    catalog_db.init()
    con = catalog_db._conn()
    con.execute(
        "INSERT INTO employees(id,name,role,model,persona) VALUES(?,?,?,?,?)",
        ("hrbp", "综合人力专员", "HRBP", "mock:model", "persona"),
    )
    con.commit()
    con.close()

    seeds.backfill_hrbp_skills()
    con = catalog_db._conn()
    assert con.execute(
        "SELECT 1 FROM employee_skills WHERE employee_id='hrbp' "
        "AND skill_id='china-unicom-ppt'").fetchone()
    assert con.execute(
        "SELECT 1 FROM employee_tools WHERE employee_id='hrbp' "
        "AND tool_id='generate_china_unicom_ppt'").fetchone()
    assert con.execute(
        "SELECT 1 FROM tools WHERE id='generate_china_unicom_ppt'").fetchone()
    con.close()
    ppt_skill = next(item for item in skill_snapshot("hrbp")
                     if item["id"] == "china-unicom-ppt")
    assert "generate_china_unicom_ppt" in ppt_skill["include_tools"]


def test_hrbp_ppt_tool_generates_editable_unicom_deck_and_preview(monkeypatch, tmp_path):
    monkeypatch.setattr(ppt_tools, "WORKSPACE_DATA", tmp_path)
    result = ppt_tools.generate_china_unicom_ppt.invoke({
        "title": "HRBP 数字员工服务能力",
        "subtitle": "员工服务场景 · 功能验证",
        "key_points": "制度答疑：先检索员工手册，再给出有依据的说明\n"
                      "流程指引：把常见人事事务整理为清晰易执行的步骤\n"
                      "闭环跟进：记录员工诉求，推动事项有序处理",
    }, config={"configurable": {"user_id": "hrbp-demo"}})

    relative = result.split("/workspace/data/", 1)[1]
    pptx_path = tmp_path / relative
    presentation = Presentation(str(pptx_path))
    assert len(presentation.slides) == 3
    text = "\n".join(shape.text for slide in presentation.slides
                     for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    assert "HRBP 数字员工服务能力" in text
    assert "制度答疑" in text
    assert "中国联通风格 PowerPoint" in result

    from app.pptx_preview import render_pptx_html
    preview = render_pptx_html(pptx_path)
    assert "制度答疑" in preview
    assert "data:image/png;base64," in preview
    assert "<script" not in preview.lower()


def test_pptx_preview_endpoint_uses_artifact_acl(monkeypatch, tmp_path):
    monkeypatch.setattr(app_paths, "WORKSPACE_DATA", tmp_path)
    pptx_path = tmp_path / "u_me" / "sample.pptx"
    pptx_path.parent.mkdir(parents=True)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 4000000, 1000000).text = "预览权限检查"
    deck.save(pptx_path)

    artifact = {
        "artifact_id": 9123, "name": "sample.pptx", "path": "u_me/sample.pptx",
        "owner_id": "u_me", "conv_id": "conv-demo", "turn_no": 1,
        "tenant_id": "default", "size": pptx_path.stat().st_size,
    }
    monkeypatch.setattr(conversations, "get_accessible_artifact",
                        lambda *args, **kwargs: artifact)

    async def keep_path(artifact_info):
        return artifact_info

    monkeypatch.setattr(workspace, "_ensure_private_snapshot", keep_path)
    app = FastAPI()
    app.include_router(workspace.router)
    user = {"id": "u_me", "username": "tester", "role": "user"}
    app.dependency_overrides[auth.get_current_user] = lambda: user
    client = TestClient(app)
    response = client.get("/api/workspace/artifacts/9123/preview")
    assert response.status_code == 200
    assert "预览权限检查" in response.text

    monkeypatch.setattr(conversations, "get_accessible_artifact",
                        lambda *args, **kwargs: None)
    denied = client.get("/api/workspace/artifacts/9123/preview")
    assert denied.status_code == 404
